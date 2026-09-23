# -*- coding: utf-8 -*-
"""`자본성증권발행현황` 시트를 지워 `sync_master_xlsx_sheet.py` 가 새로 만들게 한다.

왜 지우나: 2026-09-23 에 `법정만기일` 컬럼이 `콜만기도래일` 옆에 생겼다. sync 는 **맨 끝에
붙는 컬럼만** 자동 확장하는데(기존 컬럼이 목표의 in-order prefix 일 때만), 이 시트는 flattener
가 `비고` 를 항상 마지막에 붙이므로 새 열이 어디에 오든 중간 삽입이 되어 REFUSE 한다.
`--delete` 로 시트를 지우면 sync 의 "없는 시트 생성" 경로를 타서 목표 스키마 그대로 다시 만든다.

절차 (반드시 이 순서):
    1) python scripts/_probes/probe_20260923_capsec_sheet_recreate.py --delete
    2) python scripts/sync_master_xlsx_sheet.py "자본성증권발행현황"
    3) python scripts/_probes/probe_20260923_capsec_sheet_recreate.py --verify

가드:
  · 수식이 하나라도 있으면 중단한다(openpyxl 재저장이 캐시를 날린다 — 마스터 워크북 불변식).
  · --delete 는 **나머지 시트 전체의 값**과 대상 시트의 탭 위치를 스냅샷으로 떠 둔다.
  · --verify 는 (a) 나머지 시트가 스냅샷과 값 기준 완전 동일한지, (b) 대상 시트가 마스터 목표와
    셀 단위로 일치하는지 대조하고, (c) 탭 위치를 원래 자리로 되돌린다.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "scripts"))

from openpyxl import load_workbook  # noqa: E402

XLSX = REPO / "insurequant_master_tables.xlsx"
SHEET = "자본성증권발행현황"
SNAP = REPO / "data" / "_derived" / "_probe_20260923_capsec_sheet_snapshot.json"


def sheet_values(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def load(read_only=True):
    wb = load_workbook(XLSX, read_only=read_only, data_only=False)
    n_formula = sum(
        1 for w in wb.worksheets for row in w.iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith("="))
    if n_formula:
        sys.exit(f"REFUSE: 수식 {n_formula}개 — openpyxl 재저장은 캐시값을 날린다")
    return wb


def do_delete():
    wb = load(read_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"REFUSE: '{SHEET}' 시트가 이미 없다 — 중복 실행인지 확인해라")
    snap = {"index": wb.sheetnames.index(SHEET),
            "order": list(wb.sheetnames),
            "target_rows": len(sheet_values(wb[SHEET])),
            "others": {n: sheet_values(wb[n]) for n in wb.sheetnames if n != SHEET}}
    wb.close()
    SNAP.write_text(json.dumps(snap, ensure_ascii=False, default=str), encoding="utf-8")
    print(f"스냅샷: 시트 {len(snap['order'])}개 · '{SHEET}' 위치 {snap['index']} · "
          f"대상 {snap['target_rows']}행 · -> {SNAP.relative_to(REPO)}")

    wb = load(read_only=False)
    del wb[SHEET]
    wb.save(XLSX)
    print(f"'{SHEET}' 삭제 후 저장. 이제 sync_master_xlsx_sheet.py 를 돌려라.")


def do_verify():
    from sync_master_xlsx_sheet import target_rows
    from build_master_xlsx import MASTERS

    snap = json.loads(SNAP.read_text(encoding="utf-8"))
    wb = load(read_only=False)
    if SHEET not in wb.sheetnames:
        sys.exit(f"REFUSE: '{SHEET}' 가 없다 — sync 를 먼저 돌려라")

    # (a) 나머지 시트 불변
    bad = []
    for name, want in snap["others"].items():
        if name not in wb.sheetnames:
            bad.append(f"{name}: 시트가 사라졌다")
            continue
        got = [[str(v) if v is not None else None for v in r] for r in sheet_values(wb[name])]
        want_s = [[str(v) if v is not None else None for v in r] for r in want]
        if got != want_s:
            n = sum(1 for a, b in zip(got, want_s) if a != b)
            bad.append(f"{name}: {n}행 불일치 (got {len(got)}행 / snap {len(want_s)}행)")
    if bad:
        sys.exit("REFUSE: 다른 시트가 바뀌었다\n  " + "\n  ".join(bad))
    print(f"다른 시트 {len(snap['others'])}개 불변 확인")

    # (b) 대상 시트 == 마스터 목표
    json_file = next(j for j, s, _d in MASTERS if s == SHEET)
    cols, tgt = target_rows(json_file)
    cur = sheet_values(wb[SHEET])
    if [c for c in cur[0]] != cols:
        sys.exit(f"REFUSE: 헤더 불일치\n  시트: {cur[0]}\n  목표: {cols}")
    if len(cur) - 1 != len(tgt):
        sys.exit(f"REFUSE: 행수 {len(cur)-1} != 목표 {len(tgt)}")
    diff = [(i, j) for i, (a, b) in enumerate(zip(cur[1:], tgt))
            for j, (x, y) in enumerate(zip(a, b))
            if (None if x is None else str(x)) != (None if y is None else str(y))]
    if diff:
        for i, j in diff[:10]:
            print(f"  행{i+2} {cols[j]}: 시트={cur[i+1][j]!r} 목표={tgt[i][j]!r}")
        sys.exit(f"REFUSE: 셀 {len(diff)}개 불일치")
    print(f"'{SHEET}' {len(tgt)}행 x {len(cols)}열 마스터와 셀 단위 일치")

    # (c) 탭 위치 복원
    now = wb.sheetnames.index(SHEET)
    if now != snap["index"]:
        wb.move_sheet(SHEET, offset=snap["index"] - now)
        wb.save(XLSX)
        print(f"탭 위치 복원 {now} -> {snap['index']}")
    if wb.sheetnames != snap["order"]:
        sys.exit(f"REFUSE: 시트 순서가 다르다\n  {wb.sheetnames}\n  {snap['order']}")
    print("시트 순서 원상 복구 확인:", wb.sheetnames)


if __name__ == "__main__":
    if "--delete" in sys.argv:
        do_delete()
    elif "--verify" in sys.argv:
        do_verify()
    else:
        sys.exit("--delete 또는 --verify 를 줘라")
