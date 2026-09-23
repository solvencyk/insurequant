# -*- coding: utf-8 -*-
"""`요약` 시트의 `자본성증권발행현황` 설명 칸 1개를 build_master_xlsx.MASTERS 문구로 맞춘다.

sync_master_xlsx_sheet.py 는 요약의 **행수만** 맞추고 설명 칸은 일부러 안 건드린다(다른 레인이
손으로 고쳐 둔 문구가 있어서다). 2026-09-23 체감 기준을 법정만기로 바꾸면서 MASTERS 설명을
고쳤으므로, 그 한 칸만 손으로 반영한다. 다른 셀은 하나도 안 건드린다.
"""
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "scripts"))

from openpyxl import load_workbook  # noqa: E402
from build_master_xlsx import MASTERS  # noqa: E402

XLSX = REPO / "insurequant_master_tables.xlsx"
SHEET = "자본성증권발행현황"
want = next(d for _f, s, d in MASTERS if s == SHEET)

wb = load_workbook(XLSX, data_only=False)
n_formula = sum(1 for w in wb.worksheets for row in w.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("="))
if n_formula:
    sys.exit(f"REFUSE: 수식 {n_formula}개")
idx = wb["요약"]
hits = [r for r in range(4, idx.max_row + 1) if idx.cell(row=r, column=1).value == SHEET]
if len(hits) != 1:
    sys.exit(f"REFUSE: 요약에서 '{SHEET}' 줄이 {len(hits)}개 (1개여야 한다)")
cell = idx.cell(row=hits[0], column=4)
before = cell.value
if before == want:
    print("이미 최신 — 파일 안 씀")
    sys.exit(0)
snap = {w.title: [[c.value for c in r] for r in w.iter_rows()] for w in wb.worksheets}
cell.value = want
after = {w.title: [[c.value for c in r] for r in w.iter_rows()] for w in wb.worksheets}
diff = [(t, i, j) for t in snap for i, (a, b) in enumerate(zip(snap[t], after[t]))
        for j, (x, y) in enumerate(zip(a, b)) if x != y]
if diff != [("요약", hits[0] - 1, 3)]:
    sys.exit(f"REFUSE: 예상 밖 셀이 바뀌었다 {diff[:5]}")
wb.save(XLSX)
print(f"요약 {hits[0]}행 설명 갱신 (1셀만)\n  전: {before[:70]}...\n  후: {want[:70]}...")
