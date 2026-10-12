# -*- coding: utf-8 -*-
"""Compare two master xlsx files sheet by sheet (values, read-only). usage: z5_xlsx_sheet_diff.py A.xlsx B.xlsx"""
import sys
from openpyxl import load_workbook
sys.stdout.reconfigure(encoding="utf-8")
a, b = sys.argv[1], sys.argv[2]


def rows(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    out = {}
    for ws in wb.worksheets:
        out[ws.title] = [tuple(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return out


A, B = rows(a), rows(b)
print("sheets only in A:", sorted(set(A) - set(B)), "only in B:", sorted(set(B) - set(A)))
for name in A:
    if name in B and A[name] != B[name]:
        ra, rb = A[name], B[name]
        n = sum(1 for x, y in zip(ra, rb) if x != y) + abs(len(ra) - len(rb))
        print(f"DIFF {name}: rows {len(ra)} vs {len(rb)}, differing rows ~{n}")
