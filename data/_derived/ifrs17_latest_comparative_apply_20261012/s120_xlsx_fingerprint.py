"""Read-only: per-sheet value fingerprint (sha256 of all cell values, row-major, normalised like sync_master_xlsx_sheet.norm) of the master xlsx.
usage: s120_xlsx_fingerprint.py <xlsx> <out.json>     (the workbook is opened read-only, never saved)"""
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
src, out = Path(sys.argv[1]), Path(sys.argv[2])
wb = load_workbook(src, data_only=False, read_only=True)
res = {}
for ws in wb.worksheets:
    h = hashlib.sha256()
    n = 0
    for row in ws.iter_rows(values_only=True):
        n += 1
        h.update(("\x1f".join("" if v is None else repr(v) for v in row) + "\x1e").encode("utf-8"))
    res[ws.title] = {"rows": n, "sha256": h.hexdigest()}
wb.close()
out.write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8", newline="\n")
print(src.name, "sheets", len(res), "->", out.name)
