"""c4: row/cell change list for the three synced sheets (기본자본소진율 · 보완자본소진율 · 자본비율전망) -
backup taken before the sync vs the synced workbook, joined on (원보험사코드, 공시분기, 항목명) (read-only,
openpyxl read_only + data_only - nothing is saved).

The sync script treats 비고 as part of the row key, so a changed note shows up there as delete+insert; this list
re-joins on the business key so each change reads as old -> new for 값 / 비고 / _diagnostics.

Run (repo root):
  C:/Users/sangwook.cho/venvs/insurequant/Scripts/python.exe data/_derived/capital_utilization_apply_20261012/c4_xlsx_change_list.py
Outputs (same folder): c4_xlsx_change_list.csv (utf-8-sig) · c4_xlsx_change_summary.txt
"""
from __future__ import annotations

import csv
import io
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
BEFORE = HERE / "insurequant_master_tables.xlsx.bak_20261012_pre_capsec"
AFTER = REPO / "insurequant_master_tables.xlsx"
SHEETS = ["기본자본소진율", "보완자본소진율", "자본비율전망"]


def rows_of(path: Path, sheet: str):
    # file object, not the path: openpyxl refuses the '.bak_*' extension of the backup by name
    wb = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    out = {}
    for r in it:
        d = dict(zip(hdr, r))
        key = (d.get("원보험사코드"), d.get("공시분기"), d.get("항목명"))
        if key in out:
            raise SystemExit(f"duplicate key {key} in {path.name}:{sheet}")
        out[key] = d
    wb.close()
    return hdr, out


def norm(v):
    if v is None or v == "":
        return None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def main() -> int:
    changes, summary = [], []
    for sheet in SHEETS:
        h0, before = rows_of(BEFORE, sheet)
        h1, after = rows_of(AFTER, sheet)
        cnt = Counter()
        for key in sorted(set(before) | set(after), key=lambda k: tuple(str(x) for x in k)):
            a, b = before.get(key), after.get(key)
            if a is None or b is None:
                cnt["row_added" if a is None else "row_removed"] += 1
                changes.append({"sheet": sheet, "code": key[0], "quarter": key[1], "item": key[2],
                                "column": "(row)", "before": None if a is None else "present",
                                "after": None if b is None else "present"})
                continue
            for col in [c for c in h1 if c not in ("원보험사코드", "공시분기", "항목명")]:
                if norm(a.get(col)) != norm(b.get(col)):
                    cnt[col] += 1
                    changes.append({"sheet": sheet, "code": key[0], "quarter": key[1], "item": key[2],
                                    "column": col, "before": a.get(col), "after": b.get(col)})
        summary.append(f"{sheet}: rows before={len(before)} after={len(after)} header_same={h0 == h1} · "
                       f"changes by column: {dict(cnt)}")
    with (HERE / "c4_xlsx_change_list.csv").open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=["sheet", "code", "quarter", "item", "column", "before", "after"])
        w.writeheader()
        w.writerows(changes)
    txt = "\n".join(summary) + f"\ntotal changed cells (incl. 비고/_diagnostics): {len(changes)}\n"
    (HERE / "c4_xlsx_change_summary.txt").write_text(txt, encoding="utf-8")
    print(txt)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
