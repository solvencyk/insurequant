"""v1: prove the other sheets are value-identical before/after adding the claims-AE sheet.

Compares every sheet of insurequant_master_tables.xlsx.bak_20261012_ae_sheet (before) with the current workbook
(after), data_only/read_only (no re-save). Expected: only the new sheet is added; the summary sheet gains one row.

Run (repo root):
  C:/Users/sangwook.cho/venvs/insurequant/Scripts/python.exe data/_derived/ae_sheet_20261012/v1_other_sheets_identity.py
"""
from __future__ import annotations

import hashlib
import io
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
NEW_SHEET = "보험금예실차비율"


def sheet_hashes(path: Path) -> dict[str, tuple[str, int]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    out = {}
    for ws in wb.worksheets:
        h = hashlib.sha256()
        n = 0
        for row in ws.iter_rows(values_only=True):
            h.update(repr(row).encode("utf-8"))
            n += 1
        out[ws.title] = (h.hexdigest()[:16], n)
    wb.close()
    return out


def main() -> int:
    with tempfile.TemporaryDirectory() as td:       # openpyxl refuses the .bak extension -> read a temp copy
        tmp = Path(td) / "before.xlsx"
        shutil.copyfile(REPO / "insurequant_master_tables.xlsx.bak_20261012_ae_sheet", tmp)
        before = sheet_hashes(tmp)
    after = sheet_hashes(REPO / "insurequant_master_tables.xlsx")
    lines = [f"before {len(before)} sheets, after {len(after)} sheets"]
    added = sorted(set(after) - set(before))
    removed = sorted(set(before) - set(after))
    changed = sorted(k for k in before if k in after and before[k] != after[k])
    lines.append(f"added={added} removed={removed}")
    lines.append(f"changed (value hash differs)={changed}")
    for k in changed:
        lines.append(f"  {k}: rows {before[k][1]} -> {after[k][1]}")
    lines.append(f"unchanged sheets: {len(before) - len(changed) - len(removed)}")
    ok = added == [NEW_SHEET] and not removed and set(changed) <= {"요약"}
    lines.append("RESULT: " + ("OK" if ok else "UNEXPECTED"))
    txt = "\n".join(lines) + "\n"
    (HERE / "v1_other_sheets_identity.txt").write_text(txt, encoding="utf-8")
    print(txt)
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
