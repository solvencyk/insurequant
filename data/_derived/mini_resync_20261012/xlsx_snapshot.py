# -*- coding: utf-8 -*-
"""Snapshot every data sheet of the master xlsx (data_only, read_only) to JSON: python xlsx_snapshot.py <out.json>."""
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
out = Path(sys.argv[1])
wb = load_workbook(ROOT / "insurequant_master_tables.xlsx", data_only=True, read_only=True)
res = {}
for ws in wb.worksheets:
    rows = [json.dumps(list(r), ensure_ascii=False, default=str) for r in ws.iter_rows(values_only=True)]
    h = hashlib.sha256("\n".join(rows).encode("utf-8")).hexdigest()
    res[ws.title] = {"n": len(rows), "sha": h, "rows": rows}
out.write_text(json.dumps(res, ensure_ascii=False), encoding="utf-8")
print({k: v["n"] for k, v in res.items()})
