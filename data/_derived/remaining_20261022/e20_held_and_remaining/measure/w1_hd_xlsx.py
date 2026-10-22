"""w1_hd_xlsx: list sheets and search 보험손익 / LOB rows in the 현대해상 IR factsheet xlsx of FY2023_Q1 / FY2024_Q1 / FY2024_Q2 (READ-ONLY; openpyxl read-only, values only)."""
import sys
import openpyxl
sys.stdout.reconfigure(encoding="utf-8")
ROOT = "C:/Users/sangwook.cho/Desktop/insurequant/data/ir/"
for f in sys.argv[1:]:
    p = ROOT + f
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    print("#####", f, wb.sheetnames)
    for ws in wb.worksheets:
        n = 0
        for row in ws.iter_rows(values_only=True):
            vals = [v for v in row if v is not None]
            txt = " | ".join(str(v)[:18] for v in vals[:14])
            if any(k in txt for k in ("보험손익", "보험서비스", "자동차", "일반", "장기")) and n < 14:
                print(f"  [{ws.title}] {txt[:220]}")
                n += 1
