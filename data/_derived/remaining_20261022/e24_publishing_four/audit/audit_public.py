import json,sys,collections
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook
R='C:/Users/sangwook.cho/Desktop/insurequant/'
SP='C:/Users/sangwook.cho/AppData/Local/Temp/claude/C--Users-sangwook-cho-Desktop-insurequant/429689e1-5ae3-4467-ae2a-745748238c8b/scratchpad/'
def num(v):
    if v is None or (isinstance(v,str) and v.strip()==''): return None
    try: return float(v)
    except: return v
for t in ('CSM워터폴','손익분해PL','신계약CSM배수'):
    pe=json.load(open(R+f'public_exports/{t}.json',encoding='utf-8'))
    rows=pe if isinstance(pe,list) else pe.get('rows') or pe.get('data')
    if t=='CSM워터폴': print('public json top:',type(pe).__name__, list(pe)[:5] if isinstance(pe,dict) else json.dumps(pe[0],ensure_ascii=False)[:200])
    cols=list(rows[0]) if isinstance(rows[0],dict) else None
    def load(path):
        wb=load_workbook(path,read_only=True); ws=wb[t]; rr=[list(r) for r in ws.iter_rows(values_only=True)]; wb.close(); return rr
    out={}
    for lab,path in (('xlsx_HEAD',SP+'xlsx_head.xlsx'),('xlsx_NOW',R+'insurequant_master_tables.xlsx')):
        rr=load(path); hdr=rr[0]
        if cols is None: print('rows not dict'); break
        n=len(rr)-1
        # positional compare: same row order?
        idc=[c for c in cols if c in hdr]
        diff=0; cells=0; order_ok=len(rows)==n
        for i,r in enumerate(rr[1:]):
            d=dict(zip(hdr,r))
            if i>=len(rows): break
            for c in cols:
                cells+=1
                if num(d.get(c))!=num(rows[i].get(c)): 
                    if not(isinstance(num(d.get(c)),float) and isinstance(num(rows[i].get(c)),float) and abs(num(d.get(c))-num(rows[i].get(c)))<1e-6): diff+=1
        print(t,lab,'public rows',len(rows),'xlsx rows',n,'cols',len(cols),'cells',cells,'cell diffs public vs',lab,':',diff)
