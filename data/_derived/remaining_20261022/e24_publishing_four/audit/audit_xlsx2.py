import json, sys, subprocess, collections
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook
R='C:/Users/sangwook.cho/Desktop/insurequant/'
SP='C:/Users/sangwook.cho/AppData/Local/Temp/claude/C--Users-sangwook-cho-Desktop-insurequant/429689e1-5ae3-4467-ae2a-745748238c8b/scratchpad/'
FOUR={'KR0099','KR0003','KR0070','KR0009'}
def num(v):
    if v is None: return None
    if isinstance(v,str):
        s=v.strip()
        if s=='': return None
        try: return float(s)
        except: return s
    if isinstance(v,(int,float)): return float(v)
    return v
def same(a,b):
    a=num(a);b=num(b)
    if a is None or b is None: return a is None and b is None
    if isinstance(a,float) and isinstance(b,float): return abs(a-b)<=1e-9*max(1,abs(a),abs(b))
    return a==b
def rd(path,t):
    wb=load_workbook(path,read_only=True); ws=wb[t]
    rows=[list(r) for r in ws.iter_rows(values_only=True)]; wb.close(); return rows
def keyed(rows_iter, keyf):
    d={};cnt=collections.Counter()
    for r in rows_iter:
        k=keyf(r); c=cnt[k]; cnt[k]+=1; d[(k,c)]=r
    return d
def gitjson(rev,f):
    out=subprocess.run(['git','-C',R[:-1],'show',f'{rev}:{f}'],capture_output=True).stdout
    return json.loads(out.decode('utf-8'))
PRE=subprocess.run(['git','-C',R[:-1],'rev-parse','b1187788^'],capture_output=True,text=True).stdout.strip()
for t,jf in (('CSM워터폴','CSM_waterfall.json'),('손익분해PL','PL_breakdown.json')):
    now=rd(R+'insurequant_master_tables.xlsx',t); head=rd(SP+'xlsx_head.xlsx',t)
    hdr=now[0]
    kx=lambda r:(r[0],r[4],r[5],r[6])
    N=keyed(now[1:],kx); H=keyed(head[1:],kx)
    J=json.load(open(R+jf,encoding='utf-8'))
    JH=keyed([[r['원보험사코드'],None,None,None,r['항목번호'],r['항목명'],r['공시분기'],r.get('값'),r.get('값_당분기')] for r in J],kx)
    P=gitjson(PRE,jf)
    JP=keyed([[r['원보험사코드'],None,None,None,r['항목번호'],r['항목명'],r['공시분기'],r.get('값'),r.get('값_당분기')] for r in P],kx)
    print('\n==',t,'xlsx rows',len(N),'json rows',len(JH),'pre-load json rows',len(JP),'dups in key (json)',sum(1 for k in JH if k[1]>0))
    print(' keyset xlsx==json',set(N)==set(JH),'xlsx-json',len(set(N)-set(JH)),'json-xlsx',len(set(JH)-set(N)))
    mism=0;n=0
    for k in JH:
        if k in N:
            for ci in (7,8):
                n+=1
                if not same(N[k][ci],JH[k][ci]): mism+=1; 
    print(' cells compared',n,'mismatch xlsx vs JSON(HEAD, committed)',mism)
    # JSON change set pre->HEAD
    A=set()
    for k in set(JH)|set(JP):
        for ci in (7,8):
            a=JP.get(k,[None]*9)[ci]; b=JH.get(k,[None]*9)[ci]
            if not same(a,b): A.add((k,ci))
    B=set()
    for k in set(N)|set(H):
        for ci in (7,8):
            a=H.get(k,[None]*9)[ci]; b=N.get(k,[None]*9)[ci]
            if not same(a,b): B.add((k,ci))
    print(' JSON cells changed pre-load->HEAD:',len(A),' by co',dict(collections.Counter(k[0][0] for k,_ in A)))
    print(' xlsx cells changed HEAD->NOW:',len(B),' by co',dict(collections.Counter(k[0][0] for k,_ in B)))
    print(' A-B',len(A-B),' B-A',len(B-A))
    for x in sorted(A-B)[:10]: print('  A-B',x, 'json pre',JP.get(x[0],[None]*9)[x[1]],'json head',JH.get(x[0],[None]*9)[x[1]],'xlsx head',H.get(x[0],[None]*9)[x[1]],'xlsx now',N.get(x[0],[None]*9)[x[1]])
    for x in sorted(B-A)[:10]: print('  B-A',x, 'json pre',JP.get(x[0],[None]*9)[x[1]],'json head',JH.get(x[0],[None]*9)[x[1]],'xlsx head',H.get(x[0],[None]*9)[x[1]],'xlsx now',N.get(x[0],[None]*9)[x[1]])
    # blanked cells
    bl=[(k,ci) for (k,ci) in B if num(H[k][ci]) is not None and num(N[k][ci]) is None]
    fl=[(k,ci) for (k,ci) in B if num(H[k][ci]) is None and num(N[k][ci]) is not None]
    print(' value->blank',len(bl),' blank->value',len(fl))
    print(' sample value->blank:',[(k[0][0],k[0][1],k[0][2],k[0][3],ci,H[k][ci]) for k,ci in sorted(bl)[:12]])
    print(' value->blank by co/quarter',dict(collections.Counter((k[0][0],k[0][3]) for k,ci in bl)))
    # raw type convention
    tc=collections.Counter()
    for k,ci in B: tc[(ci,type(N[k][ci]).__name__)]+=1
    print(' NOW raw types of changed cells (col idx,type):',dict(tc))
    # for unchanged cells in col 8 what's typical type
    tt=collections.Counter(type(N[k][8]).__name__ for k in N if (k,8) not in B)
    print(' typical col8 type among unchanged:',dict(tt))
