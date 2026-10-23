import json, sys, collections
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook
R='C:/Users/sangwook.cho/Desktop/insurequant/'
wb=load_workbook(R+'insurequant_master_tables.xlsx',read_only=True); ws=wb['손익분해PL']
rows=[list(r) for r in ws.iter_rows(values_only=True)][1:]; wb.close()
J=json.load(open(R+'PL_breakdown.json',encoding='utf-8'))
X=set((r[0],r[4],r[5],r[6]) for r in rows); Jk=set((r['원보험사코드'],r['항목번호'],r['항목명'],r['공시분기']) for r in J)
a=sorted(X-Jk,key=str); b=sorted(Jk-X,key=str)
print(len(a),len(b)); print(a[:6]); print(b[:6])
print(collections.Counter(type(k[1]).__name__ for k in X), collections.Counter(type(k[1]).__name__ for k in Jk))
def num(v):
    if v is None or (isinstance(v,str) and v.strip()==''): return None
    return float(v)
xs={(r[0],r[5],r[6]):(num(r[7]),num(r[8])) for r in rows if r[4] is None}
js={(r['原보험사코드'] if False else r['원보험사코드'],r['항목명'],r['공시분기']):(num(r.get('값')),num(r.get('값_당분기'))) for r in J if isinstance(r['항목번호'],str)}
print('null-item rows xlsx',len(xs),'json str-item rows',len(js),'keys equal',set(xs)==set(js))
bad=[k for k in js if k in xs and any((a is None)!=(b is None) or (a is not None and abs(a-b)>1e-6) for a,b in zip(xs[k],js[k]))]
print('value mismatches',len(bad), 'companies',set(k[0] for k in js))
