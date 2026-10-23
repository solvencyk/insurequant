import json, sys, math, collections
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook
R='C:/Users/sangwook.cho/Desktop/insurequant/'
SP='C:/Users/sangwook.cho/AppData/Local/Temp/claude/C--Users-sangwook-cho-Desktop-insurequant/429689e1-5ae3-4467-ae2a-745748238c8b/scratchpad/'
NOW=R+'insurequant_master_tables.xlsx'; HEAD=SP+'xlsx_head.xlsx'
FOUR={'KR0099','KR0003','KR0070','KR0009'}
TARGET={'CSM워터폴':'CSM_waterfall.json','손익분해PL':'PL_breakdown.json','신계약CSM배수':'NB_CSM_multiple.json'}

def rows(path):
    wb=load_workbook(path,read_only=True,data_only=False)
    out={}
    for ws in wb.worksheets:
        out[ws.title]=[list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return out
now=rows(NOW); head=rows(HEAD)
print('sheets equal order:',list(now)==list(head), len(now))
def num(v):
    if v is None: return None
    if isinstance(v,str):
        s=v.strip()
        if s=='' : return None
        try: return float(s)
        except: return s
    if isinstance(v,bool): return v
    if isinstance(v,(int,float)): return float(v)
    return v
def same(a,b):
    a=num(a); b=num(b)
    if a is None or b is None: return a is None and b is None
    if isinstance(a,float) and isinstance(b,float):
        return abs(a-b)<=1e-9*max(1,abs(a),abs(b))
    return a==b
# (1b) other sheets identical to HEAD (strict raw compare + normalized)
for t in now:
    a=now[t]; b=head[t]
    if t in TARGET: continue
    raw=(a==b)
    nd=0
    if not raw:
        if len(a)!=len(b): nd=-1
        else:
            for x,y in zip(a,b):
                if len(x)!=len(y) or any(not same(p,q) for p,q in zip(x,y)): nd+=1
    print('OTHER',t,len(a),'raw_identical' if raw else f'DIFF normalized_diff_rows={nd}')
# types helper
def tname(v): return type(v).__name__
# (1a) target sheets: changed vs HEAD
res={}
for t,jf in TARGET.items():
    a=now[t]; b=head[t]
    hdr=a[0]; assert hdr==b[0], 'header differs'
    print('\n==',t,'rows now/head',len(a),len(b),'hdr',hdr)
    keycols=[i for i,h in enumerate(hdr) if h in('원보험사코드','항목번호','공시분기')]
    # row-aligned compare + keyed
    ka={tuple(r[i] for i in keycols):r for r in a[1:]}
    kb={tuple(r[i] for i in keycols):r for r in b[1:]}
    print(' keys equal:',set(ka)==set(kb),'dupkeys',len(a)-1-len(ka))
    print(' row order same:',[tuple(r[i] for i in keycols) for r in a[1:]]==[tuple(r[i] for i in keycols) for r in b[1:]])
    chg=collections.Counter(); outside=[]; typechg=collections.Counter(); samebut_type=0
    for k in ka:
        ra=ka[k]; rb=kb[k]
        for ci,h in enumerate(hdr):
            if ci in(1,2,3,5) : # descriptive cols
                if ra[ci]!=rb[ci]: outside.append((k,h,rb[ci],ra[ci]))
                continue
            if not same(ra[ci],rb[ci]):
                chg[(k[0],h)]+=1
                if k[0] not in FOUR: outside.append((k,h,rb[ci],ra[ci]))
                typechg[(tname(rb[ci]),tname(ra[ci]))]+=1
            elif ra[ci]!=rb[ci]:
                samebut_type+=1
    byco=collections.Counter()
    for (c,h),n in chg.items(): byco[c]+=n
    print(' changed cells by company:',dict(byco),'total',sum(byco.values()))
    print(' changed cell type transitions HEAD->NOW (non-numeric-equal):',dict(typechg))
    print(' same value but different raw (type-only) cells:',samebut_type)
    print(' changes outside FOUR/desc cols:',len(outside),outside[:5])
    res[t]=(ka,hdr,keycols)
# (1c) cell by cell vs JSON
def load(jf): return json.load(open(R+jf,encoding='utf-8'))
def cmp_sheet(t,jf):
    ka,hdr,keycols=res[t]
    J=load(jf)
    if t=='신계약CSM배수':
        keyf=lambda r:(r['원보험사코드'],r['공시분기'])
    else:
        keyf=lambda r:(r['원보험사코드'],r['항목번호'],r['공시분기'])
    kj={}
    for r in J:
        k=keyf(r); assert k not in kj,('dup json',k); kj[k]=r
    # xlsx key order is by hdr order: code, item, quarter
    if t=='신계약CSM배수': kx={ (k[0],k[1]):v for k,v in ka.items()}
    else: kx={k:v for k,v in ka.items()}
    print('\n== JSON vs xlsx',t,'json rows',len(kj),'xlsx rows',len(kx),'only json',len(set(kj)-set(kx)),'only xlsx',len(set(kx)-set(kj)))
    vcols=[h for h in hdr if h not in('원보험사코드','원수사명','티커','생손보여부','항목번호','항목명','공시분기')]
    mism=[];n=0;blank_ok=0
    descm=0
    for k in set(kj)&set(kx):
        row=dict(zip(hdr,kx[k])); jr=kj[k]
        for h in ('원수사명','생손보여부','항목명'):
            if h in row and row[h]!=jr.get(h): descm+=1
        for h in vcols:
            n+=1
            if not same(row[h],jr.get(h)): mism.append((k,h,jr.get(h),row[h]))
            elif num(jr.get(h)) is None: blank_ok+=1
    print(' value cells compared',n,'mismatch',len(mism),'blank-both',blank_ok,'desc mismatch',descm)
    for m in mism[:15]: print('  ',m)
    return mism
for t,jf in TARGET.items(): cmp_sheet(t,jf)
