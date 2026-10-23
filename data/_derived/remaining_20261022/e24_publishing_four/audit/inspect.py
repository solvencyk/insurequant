import json, sys
sys.stdout.reconfigure(encoding='utf-8')
R='C:/Users/sangwook.cho/Desktop/insurequant/'
for f in ['CSM_waterfall.json','PL_breakdown.json','NB_CSM_multiple.json']:
    d=json.load(open(R+f,encoding='utf-8'))
    print(f,type(d).__name__, list(d.keys())[:8] if isinstance(d,dict) else len(d))
    rows=d if isinstance(d,list) else None
    if rows is None:
        for k,v in d.items():
            if isinstance(v,list):
                print(' key',k,len(v)); print(json.dumps(v[0],ensure_ascii=False)[:600]); rows=v; break
            else: print(' key',k,str(v)[:200])
    else:
        print(json.dumps(rows[0],ensure_ascii=False)[:600])
from openpyxl import load_workbook
wb=load_workbook(R+'insurequant_master_tables.xlsx',read_only=True,data_only=False)
for ws in wb.worksheets:
    print(ws.title, ws.max_row, ws.max_column)
    if ws.title in('CSM워터폴','손익분해PL','신계약CSM배수'):
        for i,r in enumerate(ws.iter_rows(values_only=True)):
            print(r)
            if i>=2: break
