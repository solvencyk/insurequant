# -*- coding: utf-8 -*-
"""e23 probe (read-only): dump the '(3) UW Income' sheet of the Hyundai IR factsheet (cell coordinates + values) for the rows named in the plan."""
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
ROOT = "C:/Users/sangwook.cho/Desktop/insurequant/data/ir/"
f = sys.argv[1]
wb = openpyxl.load_workbook(ROOT + f, read_only=False, data_only=True)
print("sheets:", wb.sheetnames)
ws = wb["(3) UW Income"]
print("dims:", ws.dimensions)
lo, hi = int(sys.argv[2]), int(sys.argv[3])
for row in ws.iter_rows(min_row=lo, max_row=hi):
    cells = [(c.coordinate, c.value) for c in row if c.value is not None]
    if not cells:
        continue
    print(" | ".join(f"{k}={(f'{v:,.3f}' if isinstance(v, float) else str(v))[:22]}" for k, v in cells)[:420])
