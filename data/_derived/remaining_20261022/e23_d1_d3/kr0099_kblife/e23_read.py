# -*- coding: utf-8 -*-
"""e23 independent readers for the D1-D3 load (parser/ifrs17, ticket 20261023T0900Z).  Own code on top of the e21 table/statement helpers (e21_read / e21_raw / e21_t2); no scripts/ or src/ imports.
Contents
  read_tier1(st, role, kind, cfg)  statement items 1, 16..32 with dict citation parts (label, line, printed, sign, abs, unit, contribution_mn); cfg handles label variants (KB라이프) and expense lines
                                   printed with a minus sign (현대해상: 16 / 19 / 23 are magnitudes).
  presence_015 variants            item 15 builder default 0.0 when the INSURANCE section of the statement has no 기타영업수익 line.
  pl4_*                            CSM-amortisation (item 4) readers of the later filing's prior-period note: KB라이프 / ABL / KDB생명 / 현대해상.
  csm_*                            CSM roll-forward vectors (separate basis, issued-contract CSM) of the later filing's prior-period table: KB라이프 / ABL / KDB생명 / 현대해상 / 롯데손해보험(2023.4Q).
  hd_ir_legs(q)                    현대해상 IR factsheet '(3) UW Income' prior-year cumulative LOB insurance results (cell coordinates).
Every reader returns citations with physical xml lines (or xlsx cell coordinates) so that e23_verify.py can re-read them with its own code."""
import re
import sys
from pathlib import Path

sys.path.insert(0, "C:/Users/sangwook.cho/Desktop/insurequant/data/_derived/remaining_20261022/e21_ready25")
import e21_csm as C  # noqa: E402
import e21_raw as W  # noqa: E402
import e21_read as R  # noqa: E402
import e21_t2 as T  # noqa: E402

M = T.M
nl = T.nl
ROOT = M.ROOT
UNIT_DIV = {"원": 1e6, "천원": 1e3, "백만원": 1.0}
DIV_EOK = {"백만원": 100.0, "원": 1e8, "천원": 1e5}


class Stmt23(R.Stmt):
    """Blank-aware column picking for the plain-position layout (no ACONTEXT).  The e21 base class drops blank cells when it builds the rows, which shifts the columns of every row that prints a
    blank for a zero (KB라이프 FY2024 Q2/Q3 '영업외손익' prints a blank 3개월 cell: [blank, (362), 298, (704)] was read as [(362), 298, (704)] and 'prior cum' came out as 298 = the prior 3-MONTH
    cell).  Here every row keeps ALL value cells (blank ones included) and the column index comes from the number of value cells of the row: 4 = [cur 3m, cur cum, prior 3m, prior cum],
    2 = [cur, prior] (Q1 / annual), anything else raises.  A blank chosen cell returns None (never a guessed zero)."""

    def __init__(self, code, fy, path=None):
        super().__init__(code, fy, path)
        lines = R.X.load_lines(self.path)
        ts = R.X.find_titles(lines, r"4-2\. 포괄손익계산서")
        by_tr = {}
        if ts:
            _s, _e, rows = R.X.read_statement_at(lines, ts[0])
            for r in rows:
                by_tr[r["tr_line"]] = [dict(text=c["text"], val=c["val"], line=c["line"], ctx=c["ctx"]) for c in r["cells"]]
        else:
            s, e = R.X.html_marker_range(lines, r"4-2\. 포괄손익계산서")
            for r in R.X.read_html_rows(lines, s, e):
                if r["cells"]:
                    by_tr[r["tr_line"]] = [dict(text=c["text"], val=c["val"], line=c["line"], ctx="") for c in r["cells"][1:]]
        for row in self.rows:
            row.allcells = by_tr.get(row.tr)

    def pick_cell(self, row, role, kind):
        if any(R.RE_CTX.match(c["ctx"] or "") for c in row.cells):
            return super().pick_cell(row, role, kind)
        a = getattr(row, "allcells", None)
        if not a:
            return None
        n = len(a)
        if n == 4:
            idx = {("cur", "3m"): 0, ("cur", "cum"): 1, ("prior", "3m"): 2, ("prior", "cum"): 3}[(role, kind)]
        elif n == 2:
            if kind == "3m":
                return None
            idx = 0 if role == "cur" else 1
        else:
            raise RuntimeError(f"{self.rel}: row {row.raw!r} L{row.line}: {n} value cells (layout is not 2 or 4)")
        c = a[idx]
        if c["text"].strip() == "":
            return None
        if c["val"] is None and c["text"].strip() not in ("-", "−"):
            return None
        return c


# ============================================================================================================ Tier-1 statement reader
def mkpart(st, row, role, kind, sign=1, absf=False):
    c = st.pick_cell(row, role, kind)
    if c is None:
        raise RuntimeError(f"{st.rel}: row {row.raw!r} L{row.line}: no {role}/{kind} cell")
    v = R.cellnum(c)
    if v is None:
        raise RuntimeError(f"{st.rel}: row {row.raw!r}: cell not numeric {c['text']!r}")
    mn = v / st.scale
    return dict(label=row.raw.replace("　", " ").strip(), line=c["line"], printed=c["text"], sign=sign, abs=absf, unit=st.unit, contribution_mn=round(sign * (abs(mn) if absf else mn), 6))


def read_tier1(st, role="prior", kind="cum", cfg=None):
    """-> {item: dict(val, parts, absent?)}; helper keys '_leaf_all_sum' / '_leaf32'.  cfg: names_extra {item: [labels]}, neg_expense (bool)."""
    cfg = cfg or {}
    names = {i: list(n) for i, n in R.NAMES.items()}
    for i, extra in (cfg.get("names_extra") or {}).items():
        names[i] = names.get(i, []) + list(extra)
    neg = bool(cfg.get("neg_expense"))
    out = {}

    def direct(it, optional=False, oci_only=None, absf=False):
        r = st.find(names[it], optional=optional, oci_only=oci_only)
        if r is None:
            return None
        p = mkpart(st, r, role, kind, 1, absf)
        out[it] = dict(val=p["contribution_mn"], parts=[p])
        return out[it]

    for it in (1, 17, 20, 22, 24):
        direct(it, oci_only=False)
    # KB라이프 FY2024 Q1 layout: the insurance-finance rows form their OWN section ('보험금융손익' after the 투자영업 block), so the printed 투자영업손익 excludes them while the master's item 17
    # satisfies 20 = 1 + 17 (영업이익 = 보험영업손익 + 투자영업손익 + 보험금융손익).  Q2 / Q3 layout: the finance rows sit inside 투자영업수익 / 투자영업비용 and 투자영업손익 already includes them.
    fin_sec = st.find(["보험금융손익"], optional=True, oci_only=False)
    if fin_sec is not None:
        pf = mkpart(st, fin_sec, role, kind)
        out[17] = dict(val=out[17]["val"] + pf["contribution_mn"], parts=out[17]["parts"] + [pf])
    direct(16, oci_only=False, absf=neg)
    direct(23, oci_only=False, absf=neg)
    direct(25)
    direct(31)
    # 19 = (보험금융수익 + 재보험금융수익) - (보험금융비용 + 재보험금융비용); expense lines are magnitudes when the statement prints them with a minus sign
    parts, v19 = [], 0.0
    for nm, sg, opt in (("보험금융수익", 1, True), ("재보험금융수익", 1, True), ("보험금융비용", -1, False), ("재보험금융비용", -1, True)):
        r = st.find([nm], optional=opt, oci_only=False)
        if r is None:
            continue
        p = mkpart(st, r, role, kind, sg, neg and sg == -1)
        v19 += p["contribution_mn"]
        parts.append(p)
    out[19] = dict(val=v19, parts=parts)
    out[18] = dict(val=out[17]["val"] - v19, parts=[dict(p17) for p17 in out[17]["parts"]] + [dict(p, sign=-p["sign"], contribution_mn=round(-p["contribution_mn"], 6)) for p in parts])
    r21 = st.find(["영업외손익"], optional=True, oci_only=False)
    if r21 is not None:
        p = mkpart(st, r21, role, kind)
        out[21] = dict(val=p["contribution_mn"], parts=[p])
    else:
        a = mkpart(st, st.find(["영업외수익"], oci_only=False), role, kind)
        b = mkpart(st, st.find(["영업외비용"], oci_only=False), role, kind, -1, True)
        out[21] = dict(val=a["contribution_mn"] + b["contribution_mn"], parts=[a, b])
    cls = {26: [], 27: [], 28: [], 29: [], 30: [], 32: []}
    for r in st.leafs():
        if st.pick_cell(r, role, kind) is None:
            out.setdefault("_skipped_leafs", []).append((r.raw, r.line))
            continue
        cls[R.classify_leaf(r)].append(mkpart(st, r, role, kind))
    for it in (26, 27, 28, 29, 30):
        if len(cls[it]) > 1:
            raise RuntimeError(f"{st.rel}: item {it}: several leaf rows {[p['label'] for p in cls[it]]}")
        out[it] = dict(val=(cls[it][0]["contribution_mn"] if cls[it] else 0.0), parts=cls[it], absent=not cls[it])
    out[32] = dict(val=out[25]["val"] - sum(out[i]["val"] for i in (26, 27, 28, 29, 30)),
                   parts=[dict(out[25]["parts"][0])] + [dict(p, sign=-1, contribution_mn=round(-p["contribution_mn"], 6)) for i in (26, 27, 28, 29, 30) for p in out[i]["parts"]])
    out["_leaf32"] = cls[32]
    out["_leaf_all_sum"] = sum(p["contribution_mn"] for i in (26, 27, 28, 29, 30, 32) for p in (cls[i] if i != 32 else cls[32]))
    return out


def presence_015(st, role="prior", kind="cum", insurance_only=False):
    """item 15: the printed 기타영업수익 line of the insurance section, else the builder default 0.0.  insurance_only: ignore a 기타영업수익 row that sits after the 투자영업수익 /
    투자손익 header row (KB라이프 prints its investment-side 기타영업수익 there)."""
    inv_at = None
    for k, r in enumerate(st.rows):
        if not r.oci and r.lab in ("투자영업수익", "투자영업손익", "투자손익"):
            inv_at = k
            break
    hit = [(k, r) for k, r in enumerate(st.rows) if not r.oci and r.lab == "기타영업수익" and (not insurance_only or inv_at is None or k < inv_at)]
    if not hit:
        return dict(val=0.0, parts=[], formula="no 기타영업수익 line in the insurance section of the statement -> builder default 0.0", src=f"{st.rel}", present=False)
    if len(hit) > 1:
        raise RuntimeError(f"{st.rel}: several 기타영업수익 rows")
    p = mkpart(st, hit[0][1], role, kind)
    return dict(val=p["contribution_mn"], parts=[p], formula="printed 기타영업수익", src=f"{st.rel}", present=True)


# ============================================================================================================ small table helpers
def sep_start(lines):
    return T.sep_notes_from(lines)


CELL_OPEN = re.compile(r"^<T[DEH]\b", re.I)
CELL_CLOSE = re.compile(r"^</T[DEH]", re.I)


def rows_full23(lines, start, end):
    """like e21_raw.rows_full, but a multi-line cell also keeps the text that follows the opening tag on its FIRST line (plain <TD> dialect: '<TD ...>9. 전분기말 ...' + <SPAN> on the next line)"""
    rows, cur = [], None
    j = start - 1
    last = min(end, len(lines))
    while j < last:
        l = lines[j]
        if l.startswith("<TR"):
            cur = {"tr_line": j + 1, "cells": []}
        elif l.startswith("</TR"):
            if cur is not None and cur["cells"]:
                rows.append(cur)
            cur = None
        elif cur is not None and CELL_OPEN.match(l):
            if re.search(r"</T[DEH]>\s*$", l, re.I):
                text, ln = W.clean(l), j + 1
                nj = j
            else:
                parts, ln, k = [], None, j + 1
                t0 = W.clean(l)
                if t0:
                    parts.append(t0)
                    ln = j + 1
                while k < last and not CELL_CLOSE.match(lines[k]):
                    t = W.clean(lines[k])
                    if t:
                        parts.append(t)
                        if ln is None:
                            ln = k + 1
                    k += 1
                text = " ".join(parts).strip()
                ln = ln or (j + 1)
                nj = k
            cur["cells"].append({"text": text, "val": M.num(text), "line": ln})
            j = nj
        j += 1
    return rows


def tables_sep(path):
    """[(table, rows)] of the separate-basis notes of a filing (rows = rows_full23)"""
    lines = M.read_lines(path)
    s0 = sep_start(lines)
    out = []
    for t in M.iter_tables(lines):
        if t["start"] < s0:
            continue
        out.append((t, rows_full23(lines, t["start"], t["end"])))
    return lines, s0, out


def lab(r):
    return nl(W.label_of(r))


def eok(label, cell, sign, unit):
    return C.eok_cell(label, cell, sign, unit)


def mn_part(label, cell, unit, sign=1, absf=False):
    v, ln, tx = cell
    return dict(label=label, line=ln, printed=tx, sign=sign, abs=absf, unit=unit, contribution_mn=round(sign * (abs(v) if absf else v) / UNIT_DIV[unit], 6))


# ============================================================================================================ PL4 readers (item 4 = |CSM amortisation printed in the insurance-revenue note|), later filing, prior-period cumulative
def pl4_kblife(q):
    """KB라이프: separate notes, table '(2) 당분기/당반기 및 전분기/전반기 중 기타사업비용을 제외한 보험영업수익 및 보험영업비용' whose header names the PRIOR period (전분기/전반기), row
    '서비스 제공에 따른 보험계약마진의 변동' (first one = 보험수익 section); the 합계 of the cumulative block = the LAST numeric cell of the row."""
    fy = M.later_fy(q)
    path = T.main_path("KR0099", fy)
    lines, s0, tabs = tables_sep(path)
    hits = []
    for t, rows in tabs:
        cap = nl(t["caption"] or "")
        if "기타사업비용을제외한" not in cap or "보험영업수익" not in cap:
            continue
        hdr = nl(" ".join(c["text"] for r in rows[:3] for c in r["cells"]))
        if not re.search(r"구분(전분기|전반기|전기)", hdr):
            continue
        rr = [r for r in rows if lab(r) == "서비스제공에따른보험계약마진의변동"]
        if rr:
            hits.append((t, rr[0]))
    if len(hits) != 1:
        raise RuntimeError(f"{path.name}: {len(hits)} KB라이프 prior 보험영업수익 note tables (need 1)")
    t, r = hits[0]
    unit = T.find_unit(lines, t, span=2500)
    n = W.numeric_cells(r)
    return dict(part=mn_part("서비스 제공에 따른 보험계약마진의 변동 [합계, 누적]", n[-1], unit, 1, True), n_numeric=len(n), src=f"{M.rel(path)} table@L{t['start']} unit {unit}", unit=unit)


def pl4_abl(q):
    """ABL: separate notes, table whose first section header is 보험수익 (not the reinsurance / cost tables), row '보험계약마진상각'; prior cumulative = the LAST numeric cell."""
    fy = M.later_fy(q)
    path = T.main_path("KR0070", fy)
    lines, s0, tabs = tables_sep(path)
    hits = []
    for t, rows in tabs:
        labs = [lab(r) for r in rows]
        if "보험계약마진상각" not in labs or "예상보험금" not in labs:
            continue
        # the revenue table starts with the section header '보험수익' and has no 보험서비스비용 section header
        if "보험수익" not in labs[:6] or "보험서비스비용" in labs:
            continue
        hdr = nl(" ".join(c["text"] for r in rows[:3] for c in r["cells"]))
        if not re.search(r"구분(당분기|당반기|당기)(전분기|전반기|전기)|구분(당분기|당반기|당기)", hdr):
            continue
        hits.append((t, [r for r in rows if lab(r) == "보험계약마진상각"][0]))
    if len(hits) != 1:
        raise RuntimeError(f"{path.name}: {len(hits)} ABL 보험수익 note tables with a 보험계약마진상각 row (need 1): {[h[0]['start'] for h in hits]}")
    t, r = hits[0]
    unit = T.find_unit(lines, t, span=2500)
    n = W.numeric_cells(r)
    return dict(part=mn_part("보험계약마진상각 [전기/전분기 누적]", n[-1], unit, 1, True), n_numeric=len(n), src=f"{M.rel(path)} table@L{t['start']} unit {unit}", unit=unit)


def pl4_kdb(q):
    """KDB생명: separate notes, '보험영업수익 및 재보험영업수익의 구성내역' table, row '(3) 제공된 서비스의 보험계약마진' (the revenue side); prior cumulative = the LAST numeric cell."""
    fy = M.later_fy(q)
    path = T.main_path("KR0072", fy)
    lines, s0, tabs = tables_sep(path)
    hits = []
    for t, rows in tabs:
        cap = nl(t["caption"] or "")
        if "보험영업수익" not in cap or "구성내역" not in cap or "재보험영업수익" not in cap:
            continue
        rr = [r for r in rows if re.match(r"^\(3\)제공된서비스의보험계약마진$", lab(r))]
        if rr:
            hits.append((t, rr[0]))
    if len(hits) != 1:
        raise RuntimeError(f"{path.name}: {len(hits)} KDB 보험영업수익 구성내역 tables with the 제공된 서비스의 보험계약마진 row (need 1)")
    t, r = hits[0]
    unit = T.find_unit(lines, t, span=2500)
    n = W.numeric_cells(r)
    return dict(part=mn_part("(3) 제공된 서비스의 보험계약마진 [전기/전분기 누적]", n[-1], unit, 1, True), n_numeric=len(n), src=f"{M.rel(path)} table@L{t['start']} unit {unit}", unit=unit)


HD_PRIOR_CAP = {"1": r"^\(2\)전분기$", "2": r"^\(3\)전반기누적$", "3": r"^\(3\)전분기누적$", "4": r"^\(2\)전기$"}


def pl4_hd(q):
    """현대해상: separate notes, table '구 분 장기 자동차 일반 합계' with section 보험수익; caption = the prior-period CUMULATIVE table of the later filing; row '보험계약마진상각' (first = 보험수익 section);
    item 4 = the 장기 column (first numeric cell)."""
    fy = M.later_fy(q)
    path = T.main_path("KR0009", fy)
    lines, s0, tabs = tables_sep(path)
    hits = []
    for t, rows in tabs:
        cap = nl(t["caption"] or "")
        if not re.match(HD_PRIOR_CAP[q[5]], cap):
            continue
        hdr = nl(" ".join(c["text"] for r in rows[:3] for c in r["cells"]))
        if not hdr.startswith("구분장기자동차일반합계"):
            continue
        rr = [r for r in rows if lab(r) == "보험계약마진상각"]
        if rr:
            hits.append((t, rr[0]))
    if len(hits) != 1:
        raise RuntimeError(f"{path.name}: {len(hits)} 현대해상 prior 보험수익 note tables (need 1): {[h[0]['start'] for h in hits]}")
    t, r = hits[0]
    unit = T.find_unit(lines, t, span=2500)
    n = W.numeric_cells(r)
    return dict(part=mn_part("보험계약마진상각 [장기, 전기/전분기 누적]", n[0], unit, 1, True), n_numeric=len(n), src=f"{M.rel(path)} table@L{t['start']} unit {unit}", unit=unit)


# ============================================================================================================ CSM roll-forward readers
def finish(cells, src, derived=None):
    """cells {item: [eok_cell...]} -> contract of e21_csm.finish (vec_exact / vec 2dp / vec1 1dp)"""
    return C.finish(cells, False, src)


def _csm_rows(rows, rules, col, unit, path_name, sub_rows=None):
    """rules {item: predicate(label_norm)->bool}; cell = numeric cell number `col` (0-based among numeric cells after the label) of the UNIQUE row (first match when nth given)"""
    cells = {}
    for it, (pred, nth, label) in rules.items():
        h = [r for r in rows if pred(lab(r))]
        if len(h) < nth:
            raise RuntimeError(f"{path_name}: item {it}: {len(h)} rows match (need >= {nth})")
        r = h[nth - 1]
        n = W.numeric_cells(r)
        if len(n) <= col:
            raise RuntimeError(f"{path_name}: item {it} row {W.label_of(r)!r}: {len(n)} numeric cells (< {col + 1})")
        cells[it] = [eok(W.label_of(r), n[col], 1, unit)]
    return cells


def _unique_prior_table(tabs, pred, path_name, expect=1):
    hits = [(t, rows) for t, rows in tabs if pred(t, rows)]
    if len(hits) != expect:
        raise RuntimeError(f"{path_name}: {len(hits)} candidate CSM tables (need {expect}): {[h[0]['start'] for h in hits]}")
    return hits


def csm_kblife(q):
    """KB라이프 FY+1 separate notes: '(2) 당분기와 전분기의 보험료배분접근법을 적용하지 않은 보험계약부채(자산)의 측정요소 별 변동내역' prior (전분기/전반기/전기) table, CSM = 3rd numeric column."""
    fy = M.later_fy(q)
    path = T.main_path("KR0099", fy)
    lines, s0, tabs = tables_sep(path)

    def pred(t, rows):
        cap = nl(t["caption"] or "")
        if "보험료배분접근법을적용하지않은보험계약부채(자산)의측정요소별변동내역" not in cap:
            return False
        hdr = nl(" ".join(c["text"] for r in rows[:3] for c in r["cells"]))
        labs = [lab(r) for r in rows]
        return bool(re.search(r"구분(전분기|전반기|전기)", hdr)) and "신계약효과" in labs and any(x.startswith("보험금융손익(PL)") for x in labs)
    t, rows = _unique_prior_table(tabs, pred, path.name)[0]
    unit = T.find_unit(lines, t, span=2500)
    rules = {
        1: (lambda x: x == "기초순장부금액", 1, ""),
        2: (lambda x: x == "신계약효과", 1, ""),
        3: (lambda x: x == "보험금융손익(PL)", 1, ""),
        4: (lambda x: x == "보험계약마진을조정하는추정치변동", 1, ""),
        5: (lambda x: x == "서비스제공에따른보험계약마진변동", 1, ""),
        6: (lambda x: re.match(r"^(분기말|반기말|기말)순장부금액$", x) is not None, 1, ""),
    }
    cells = _csm_rows(rows, rules, 2, unit, path.name)
    grp = [r for r in rows if lab(r).startswith("당기서비스관련변동")]
    out = finish(cells, f"{M.rel(path)} table@L{t['start']} unit {unit}")
    out["check_group_row_5"] = eok("당기서비스관련 변동 [CSM]", W.numeric_cells(grp[0])[2], 1, unit)
    return out


def csm_abl(q):
    """ABL FY+1 separate notes: '2) 전분기|전반기|전기' roll-forward table [최선추정, 위험조정, 보험계약마진, 합계]; CSM = 3rd numeric column."""
    fy = M.later_fy(q)
    path = T.main_path("KR0070", fy)
    lines, s0, tabs = tables_sep(path)

    def pred(t, rows):
        cap = nl(t["caption"] or "")
        if not re.match(r"^2\)(전분기|전반기|전기)$", cap):
            return False
        labs = [lab(r) for r in rows]
        return any(x.startswith("I.기초보험계약순부채") for x in labs) and any("처음인식한계약의효과" in x for x in labs)
    t, rows = _unique_prior_table(tabs, pred, path.name)[0]
    unit = T.find_unit(lines, t, span=2500)
    rules = {
        1: (lambda x: x.startswith("I.기초보험계약순부채"), 1, ""),
        2: (lambda x: x.startswith("(3)해당기간에처음인식한계약의효과"), 1, ""),
        3: (lambda x: x.startswith("VI.보험금융비용"), 1, ""),
        4: (lambda x: x.startswith("(1)보험계약마진을조정하는추정치의변동분"), 1, ""),
        5: (lambda x: x.startswith("(1)서비스제공으로당기손익으로인식된보험계약마진"), 1, ""),
        6: (lambda x: re.match(r"^X\.(분기말|반기말|기말)보험계약순부채", x) is not None, 1, ""),
    }
    cells = _csm_rows(rows, rules, 2, unit, path.name)
    grp = [r for r in rows if lab(r).startswith("III.현재서비스와관련된변동분")]
    out = finish(cells, f"{M.rel(path)} table@L{t['start']} unit {unit}")
    out["check_group_row_5"] = eok("III. 현재 서비스와 관련된 변동분 [CSM]", W.numeric_cells(grp[0])[2], 1, unit)
    return out


def csm_kdb(q, role="prior"):
    """KDB생명 FY+1 separate notes: '2) 제36(전)기 N분기|반기|' caption roll-forward table [PV, RA, CSM, 합계]; CSM = 3rd numeric column.
    role='cur': the ORIGINAL filing of the quarter itself (FY of q), caption '1) 제36(당)기 ...' (used to classify the owner-gold cells that equal the original print)."""
    fy = M.later_fy(q) if role == "prior" else M.fy_of(q)
    path = T.main_path("KR0072", fy)
    lines, s0, tabs = tables_sep(path)
    cap_re = r"^2\)제36\(전\)기" if role == "prior" else r"^1\)제36\(당\)기"

    def pred(t, rows):
        cap = nl(t["caption"] or "")
        if not re.match(cap_re, cap):
            return False
        labs = [lab(r) for r in rows]
        return any(x.startswith("1.기초보험계약자산및부채") for x in labs) and any("최초인식한계약의효과" in x for x in labs) and not any("재보험계약" in x for x in labs[:6])
    t, rows = _unique_prior_table(tabs, pred, path.name)[0]
    unit = T.find_unit(lines, t, span=2500)
    rules = {
        1: (lambda x: x.startswith("1.기초보험계약자산및부채"), 1, ""),
        2: (lambda x: x.startswith("(3)최초인식한계약의효과"), 1, ""),
        3: (lambda x: x.startswith("6.보험금융손익"), 1, ""),
        4: (lambda x: x.startswith("(1)보험계약마진을조정하는추정치의변동분"), 1, ""),
        5: (lambda x: x.startswith("(1)서비스이전을반영하기위해당기손익으로인식한보험계약마진"), 1, ""),
        6: (lambda x: re.match(r"^9\.(전분기말|전반기말|전기말|분기말|반기말|기말)보험계약자산및부채", x) is not None, 1, ""),
    }
    cells = _csm_rows(rows, rules, 2, unit, path.name)
    grp = [r for r in rows if lab(r).startswith("3.현행서비스관련변동")]
    out = finish(cells, f"{M.rel(path)} table@L{t['start']} unit {unit}")
    out["check_group_row_5"] = eok("3. 현행 서비스 관련 변동 [CSM]", W.numeric_cells(grp[0])[2], 1, unit)
    return out


def csm_hd(q):
    """현대해상 FY+1 separate notes: the pair [당기, 전기] of the '(2) 최선추정, 위험조정, 보험계약마진 변동에 대한 차이조정 1) 장기원수' roll-forward tables (header: 구분 | BEL | RA | 보험계약마진 | 합계;
    rows 기초순보험계약부채 ... 신계약효과 ... 기말|당기말보험계약순부채); the SECOND table of the pair is the prior period (its caption is '나. 전...' or, in the half-year filing, the stale
    '가. 당반기' caption of the first).  CSM = 3rd numeric column."""
    fy = M.later_fy(q)
    path = T.main_path("KR0009", fy)
    lines, s0, tabs = tables_sep(path)

    def pred(t, rows):
        hdr = nl(" ".join(c["text"] for r in rows[:2] for c in r["cells"]))
        labs = [lab(r) for r in rows]
        return ("보험계약마진합계" in hdr and "기초순보험계약부채" in labs and "신계약효과" in labs and ("기말보험계약순부채" in labs or "당기말보험계약순부채" in labs)
                and "보험계약마진을조정하는변동" in labs and "보험계약마진상각" in labs)
    hits = _unique_prior_table(tabs, pred, path.name, expect=2)
    (t0, rows0), (t, rows) = hits
    cap = nl(t["caption"] or "")
    if not (re.match(r"^나\.(전분기|전반기|전기)$", cap) or (q[5] == "2" and "1)장기원수가.당반기" in cap)):
        raise RuntimeError(f"{path.name}: the second table of the pair has caption {cap!r}")
    unit = T.find_unit(lines, t, span=2500)
    rules = {
        1: (lambda x: x == "기초순보험계약부채", 1, ""),
        2: (lambda x: x == "신계약효과", 1, ""),
        3: (lambda x: x == "보험금융손익", 1, ""),
        4: (lambda x: x == "보험계약마진을조정하는변동", 1, ""),
        5: (lambda x: x == "보험계약마진상각", 1, ""),
        6: (lambda x: x in ("기말보험계약순부채", "당기말보험계약순부채"), 1, ""),
    }
    cells = _csm_rows(rows, rules, 2, unit, path.name)
    grp = [r for r in rows if lab(r) == "당기서비스와관련된변동"]
    out = finish(cells, f"{M.rel(path)} table@L{t['start']} (pair L{t0['start']}/L{t['start']}) unit {unit}")
    out["check_group_row_5"] = eok("당기서비스와 관련된 변동 [CSM]", W.numeric_cells(grp[0])[2], 1, unit)
    return out


# ============================================================================================================ 현대해상 IR factsheet legs
IR_FILES = {"2023.1Q": "data/ir/FY2024_Q1/raw/KR0009_현대해상/dd7e547e6e1f9cc171b77dcb4bf67d58.xlsx", "2023.2Q": "data/ir/FY2024_Q2/raw/KR0009_현대해상/d33c3d385303361798f75607750ca360.xlsx"}


def hd_ir_legs(q):
    """-> dict(file, sheet, header, unit, cells={'일반': {...}, '장기': ..., '자동차': ..., '합계': ...}) from '(3) UW Income' of the IR factsheet of the LATER year (prior-year cumulative column)."""
    import openpyxl
    rel = IR_FILES[q]
    wb = openpyxl.load_workbook(ROOT / rel, data_only=True)
    ws = wb["(3) UW Income"]
    # locate the prior-year cumulative column: header row 5/6 text '23.1Q' with '누계' in the same column (rows 6)
    hdr_col = None
    for c in ws[6]:
        if isinstance(c.value, str) and c.value.replace(" ", "").replace("\n", "").startswith(f"{q[2:4]}.{q[5]}Q누계"):
            hdr_col = c.column
    if hdr_col is None:
        raise RuntimeError(f"{rel}: prior-year cumulative header '{q[2:4]}.{q[5]}Q 누계' not found in row 6")
    unit_txt = None
    for c in ws[4]:
        if isinstance(c.value, str) and "단위" in c.value:
            unit_txt = c.value
    blk, res = None, {}
    for row in ws.iter_rows(min_row=8):
        a = [c for c in row if c.column <= 2 and isinstance(c.value, str)]
        for c in a:
            v = c.value.replace(" ", "")
            if v in ("일반", "장기", "자동차"):
                blk = v
        lab_cells = [c for c in row if isinstance(c.value, str) and c.value.replace(" ", "") in ("보험손익", "보험손익합계")]
        if not lab_cells:
            continue
        key = "합계" if lab_cells[0].value.replace(" ", "") == "보험손익합계" else blk
        cell = ws.cell(row=row[0].row, column=hdr_col)
        if key in res or not isinstance(cell.value, (int, float)):
            continue
        res[key] = dict(cell=cell.coordinate, row=cell.row, value_mn=cell.value, label=lab_cells[0].value, label_cell=lab_cells[0].coordinate, block=blk)
    return dict(file=rel, sheet="(3) UW Income", header_cell=ws.cell(row=6, column=hdr_col).coordinate, header_text=ws.cell(row=6, column=hdr_col).value, unit=unit_txt, cells=res)
