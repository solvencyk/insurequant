# -*- coding: utf-8 -*-
"""1c: PL sheet vs HEAD master ignoring only the 항목번호 column (pre-existing blank for Koreanre '2-1' style ids)."""
import json, subprocess, sys
from collections import Counter
from openpyxl import load_workbook
sys.path.insert(0, "C:/Users/sangwook.cho/Desktop/insurequant/scripts")
sys.stdout.reconfigure(encoding="utf-8")
exec(open("pv2_xlsx2.py", encoding="utf-8").read().split("from build_master_xlsx")[0].split("def git_json")[0])
REPO = "C:/Users/sangwook.cho/Desktop/insurequant"
out = subprocess.run(["git", "show", "HEAD:PL_breakdown.json"], cwd=REPO, capture_output=True, check=True).stdout
m = json.loads(out.decode("utf-8"))
wb = load_workbook(X, read_only=True)
rows = [list(r) for r in wb["손익분해PL"].iter_rows(values_only=True)]
hdr = rows[0]
ign = hdr.index("항목번호")
xs = Counter(tuple(n(c) for k, c in enumerate(r) if k != ign) for r in rows[1:] if any(c is not None for c in r))
ms = Counter(tuple(n(r.get(h)) for k, h in enumerate(hdr) if k != ign) for r in m)
print("only_in_xlsx", sum((xs - ms).values()), "only_in_master", sum((ms - xs).values()))
blank = [r for r in rows[1:] if r[ign] is None]
nonint = [r for r in m if not isinstance(r["항목번호"], int)]
print("xlsx blank 항목번호 rows:", len(blank), "companies", Counter(r[0] for r in blank), "master non-int ids:", len(nonint), Counter(r["원보험사코드"] for r in nonint))
