# -*- coding: utf-8 -*-
"""Independent check 1b: whole-sheet multiset equality of working xlsx vs committed master (HEAD) for the 3 sheets."""
import json
import re
import subprocess
import sys
from collections import Counter
from openpyxl import load_workbook

REPO = "C:/Users/sangwook.cho/Desktop/insurequant"
sys.path.insert(0, REPO + "/scripts")
sys.stdout.reconfigure(encoding="utf-8")
X = REPO + "/insurequant_master_tables.xlsx"
SP = "C:/Users/sangwook.cho/AppData/Local/Temp/claude/C--Users-sangwook-cho-Desktop-insurequant/429689e1-5ae3-4467-ae2a-745748238c8b/scratchpad"
XH = SP + "/xlsx_head_pv2.xlsx"
NUM = re.compile(r"^-?\d+(\.\d+)?([eE][-+]?\d+)?$")


def n(v):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, str):
        if NUM.match(v):
            v = float(v)
        else:
            return v
    f = float(v)
    return int(f) if f.is_integer() else float("%.15g" % f)


def git_json(rev, path):
    out = subprocess.run(["git", "show", f"{rev}:{path}"], cwd=REPO, capture_output=True, check=True).stdout
    return json.loads(out.decode("utf-8"))


from build_master_xlsx import FLATTEN  # read-only helper for ILP label columns

for sheet, mj in (("17BS", "IFRS17_BS.json"), ("손익분해PL", "PL_breakdown.json"), ("보험계약부채", "insurance_liability_portfolio.json")):
    m = git_json("HEAD", mj)
    if mj in FLATTEN:
        m = FLATTEN[mj](m)
    for path, tag in ((X, "work"), (XH, "headcopy")):
        wb = load_workbook(path, read_only=True, data_only=False)
        rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
        wb.close()
        hdr = rows[0]
        xs = Counter(tuple(n(c) for c in r) for r in rows[1:] if any(c is not None for c in r))
        ms = Counter(tuple(n(r.get(h)) for h in hdr) for r in m)
        only_x = xs - ms
        only_m = ms - xs
        print(f"{sheet} [{tag}] vs HEAD master: sheet rows={sum(xs.values())} master rows={sum(ms.values())} only_in_xlsx={sum(only_x.values())} only_in_master={sum(only_m.values())}")
        if tag == "work" and (only_x or only_m):
            for t in list(only_x)[:5]:
                print("   XLSX-only", t)
            for t in list(only_m)[:5]:
                print("   MASTER-only", t)
    # type consistency of 값_당분기 in PL
    if sheet == "손익분해PL":
        wb = load_workbook(X, read_only=True)
        rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
        wb.close()
        j = rows[0].index("값_당분기")
        print("  값_당분기 cell types in work xlsx:", Counter(type(r[j]).__name__ for r in rows[1:]))
        wb = load_workbook(XH, read_only=True)
        rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
        wb.close()
        print("  값_당분기 cell types in head xlsx:", Counter(type(r[j]).__name__ for r in rows[1:]))
