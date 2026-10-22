# -*- coding: utf-8 -*-
"""Peek: sheet headers (working xlsx), master row shapes, formula tag counts. Read-only."""
import json
import re
import subprocess
import sys
import zipfile
from openpyxl import load_workbook

REPO = "C:/Users/sangwook.cho/Desktop/insurequant"
X = REPO + "/insurequant_master_tables.xlsx"
SP = "C:/Users/sangwook.cho/AppData/Local/Temp/claude/C--Users-sangwook-cho-Desktop-insurequant/429689e1-5ae3-4467-ae2a-745748238c8b/scratchpad"
XH = SP + "/xlsx_head_pv2.xlsx"
sys.stdout.reconfigure(encoding="utf-8")

for p in (XH, X):
    with zipfile.ZipFile(p) as z:
        names = z.namelist()
        nf = 0
        for n in names:
            if n.startswith("xl/worksheets/") and n.endswith(".xml"):
                data = z.read(n)
                nf += len(re.findall(rb"<f[\s/>]", data))
        risky = [n for n in names if any(k in n for k in ("pivot", "chart", "drawing", "media", "/table", "vml", "queryTable", "connections"))]
        print(p.split("/")[-1], "parts", len(names), "formula tags", nf, "risky parts", risky)

wb = load_workbook(X, read_only=True)
for ws in wb.worksheets:
    hdr = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    print(ws.title, ws.max_row, list(hdr))

for fn in ("PL_breakdown.json", "IFRS17_BS.json", "insurance_liability_portfolio.json"):
    d = json.load(open(REPO + "/" + fn, encoding="utf-8"))
    print(fn, type(d).__name__, len(d))
    rows = d if isinstance(d, list) else None
    if rows:
        print(rows[0])
