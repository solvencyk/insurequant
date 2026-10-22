# -*- coding: utf-8 -*-
"""Independent check 1: xlsx (working) vs xlsx (git HEAD copy) vs committed masters. Read-only."""
import json
import subprocess
import sys
from collections import Counter, defaultdict
from openpyxl import load_workbook

REPO = "C:/Users/sangwook.cho/Desktop/insurequant"
X = REPO + "/insurequant_master_tables.xlsx"
SP = "C:/Users/sangwook.cho/AppData/Local/Temp/claude/C--Users-sangwook-cho-Desktop-insurequant/429689e1-5ae3-4467-ae2a-745748238c8b/scratchpad"
XH = SP + "/xlsx_head_pv2.xlsx"
sys.stdout.reconfigure(encoding="utf-8")


def git_json(rev, path):
    out = subprocess.run(["git", "show", f"{rev}:{path}"], cwd=REPO, capture_output=True, check=True).stdout
    return json.loads(out.decode("utf-8"))


def nrm(v):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        f = float(v)
        return int(f) if f.is_integer() else float("%.15g" % f)
    if isinstance(v, str) and v.lstrip("-").isdigit() and str(int(v)) == v:
        return int(v)
    return v


def read_all(path, data_only):
    wb = load_workbook(path, read_only=True, data_only=data_only)
    out = {}
    for ws in wb.worksheets:
        out[ws.title] = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return out


res = {}
for mode in (False, True):
    a = read_all(XH, mode)
    b = read_all(X, mode)
    print(f"[data_only={mode}] sheet list identical:", list(a) == list(b), len(a))
    changed_sheets = {}
    for name in a:
        ra, rb = a[name], b[name]
        # trim trailing all-None cols/rows effects: compare normalised padded
        n = max(len(ra), len(rb))
        diffs = 0
        for i in range(n):
            x = ra[i] if i < len(ra) else []
            y = rb[i] if i < len(rb) else []
            w = max(len(x), len(y))
            for j in range(w):
                u = x[j] if j < len(x) else None
                v = y[j] if j < len(y) else None
                if u != v:
                    diffs += 1
        if diffs or len(ra) != len(rb):
            changed_sheets[name] = (len(ra), len(rb), diffs)
    print(f"[data_only={mode}] sheets with any raw-cell difference (rows_old, rows_new, positional-diff cells):", changed_sheets)
    if not mode:
        res["a"], res["b"] = a, b
        # formulas? any str starting with '=' in either workbook
        nform = sum(1 for wbk in (a, b) for rows in wbk.values() for r in rows for c in r if isinstance(c, str) and c.startswith("="))
        print("cells that are formula strings (either workbook):", nform)

A, B = res["a"], res["b"]

# ---- keyed diff for the three sheets
SPEC = {
    "17BS": ("IFRS17_BS.json", "59e6e459^"),
    "손익분해PL": ("PL_breakdown.json", "e01f1395^"),
    "보험계약부채": ("insurance_liability_portfolio.json", "59e6e459^"),
}


def keyed(rows):
    hdr = rows[0]
    ic, iq, ii = hdr.index("원보험사코드"), hdr.index("공시분기"), hdr.index("항목번호")
    d = defaultdict(list)
    for r in rows[1:]:
        if all(c is None for c in r):
            continue
        d[(r[ic], r[iq], nrm(r[ii]))].append(dict(zip(hdr, r)))
    return hdr, d


xl_changes = {}
for sheet, (mjson, prev) in SPEC.items():
    ha, da = keyed(A[sheet])
    hb, db = keyed(B[sheet])
    print(f"\n== {sheet}: header identical={ha == hb}; rows old/new = {sum(len(v) for v in da.values())}/{sum(len(v) for v in db.values())};"
          f" dup keys old/new = {sum(1 for v in da.values() if len(v)>1)}/{sum(1 for v in db.values() if len(v)>1)}")
    cells = set()
    for k in set(da) | set(db):
        ra, rb = da.get(k, [None]), db.get(k, [None])
        if len(ra) != len(rb):
            print("  key cardinality changed", k, len(ra), len(rb))
        for x, y in zip(ra, rb):
            if x is None or y is None:
                print("  key added/removed", k)
                continue
            for c in hb:
                if nrm(x.get(c)) != nrm(y.get(c)):
                    cells.add((k, c, nrm(x.get(c)), nrm(y.get(c))))
    xl_changes[sheet] = cells
    print(f"  changed cells (keyed): {len(cells)}; companies={Counter(k[0][0] for k in cells)}; quarters={dict(Counter(k[0][1] for k in cells))}")
    # master changes in the commit(s) under check
    new = git_json("HEAD", mjson)
    old = git_json(prev, mjson)
    def mk(rows):
        d = defaultdict(list)
        for r in rows:
            d[(r["원보험사코드"], r["공시분기"], nrm(r["항목번호"]))].append(r)
        return d
    mn, mo = mk(new), mk(old)
    mcells = set()
    for k in set(mn) | set(mo):
        ra, rb = mo.get(k, [{}]), mn.get(k, [{}])
        if len(ra) != 1 or len(rb) != 1:
            print("  master dup/absent key", k, len(ra), len(rb))
        for c in set(ra[0]) | set(rb[0]):
            if c in ("원보험사코드", "공시분기", "항목번호"):
                continue
            if nrm(ra[0].get(c)) != nrm(rb[0].get(c)):
                mcells.add((k, c, nrm(ra[0].get(c)), nrm(rb[0].get(c))))
    print(f"  master cells changed {prev} -> HEAD: {len(mcells)}; keys={sorted({m[0] for m in mcells})[:40]}")
    # compare for the value columns (sheet columns 값/값_당분기); ILP also flat labels
    sheet_cols = {c for (_k, c, _o, _n) in cells}
    mcells_v = {(k, c, o, n) for (k, c, o, n) in mcells if c in ("값", "값_당분기")}
    xcells_v = {(k, c, o, n) for (k, c, o, n) in cells if c in ("값", "값_당분기")}
    # normalise None vs missing
    print(f"  value-cells: xlsx {len(xcells_v)} vs master {len(mcells_v)}; equal sets: {xcells_v == mcells_v}")
    if xcells_v != mcells_v:
        for m in sorted(xcells_v ^ mcells_v, key=str)[:40]:
            print("    XOR", m, "in xlsx" if m in xcells_v else "in master")
    other = {c for (_k, c, _o, _n) in cells} - {"값", "값_당분기"}
    print("  non-value columns changed in xlsx:", other)
    if sheet == "손익분해PL":
        pass
    for m in sorted(cells, key=str)[:60]:
        print("   xl", m)
print("\nDONE")
