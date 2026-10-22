# -*- coding: utf-8 -*-
"""pv3 check 1: working xlsx vs git HEAD xlsx copy (all sheets, cell level) + changed PL cells vs HEAD master
and vs the pre-change master (HEAD~1). Read-only (openpyxl read_only, never saves)."""
import json
import re
import subprocess
import sys
from collections import Counter
from openpyxl import load_workbook

REPO = "C:/Users/sangwook.cho/Desktop/insurequant"
SP = "C:/Users/sangwook.cho/AppData/Local/Temp/claude/C--Users-sangwook-cho-Desktop-insurequant/429689e1-5ae3-4467-ae2a-745748238c8b/scratchpad"
X = REPO + "/insurequant_master_tables.xlsx"
XH = SP + "/xlsx_head_pv3.xlsx"
sys.stdout.reconfigure(encoding="utf-8")
NUM = re.compile(r"^-?\d+(\.\d+)?([eE][-+]?\d+)?$")


def n(v):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, str):
        if NUM.match(v):
            v = float(v)
        else:
            return v
    f = float(v)
    return int(f) if f.is_integer() else float("%.15g" % f)


def gj(rev, path):
    out = subprocess.run(["git", "show", f"{rev}:{path}"], cwd=REPO, capture_output=True, check=True).stdout
    return json.loads(out.decode("utf-8"))


def sheets(path):
    wb = load_workbook(path, read_only=True, data_only=False)
    d = {}
    for ws in wb.worksheets:
        d[ws.title] = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return d


W, H = sheets(X), sheets(XH)
print("sheet names equal:", list(W) == list(H), len(W))
formulas = 0
for s, rows in W.items():
    for r in rows:
        for c in r:
            if isinstance(c, str) and c.startswith("="):
                formulas += 1
print("formula-like cells in working xlsx:", formulas)
changed = {}
for s in W:
    a, b = H[s], W[s]
    if len(a) != len(b):
        print(s, "ROW COUNT differs", len(a), len(b))
        continue
    ch = []
    for i, (ra, rb) in enumerate(zip(a, b)):
        if len(ra) != len(rb):
            ch.append((i, "LEN", len(ra), len(rb)))
            continue
        for j, (ca, cb) in enumerate(zip(ra, rb)):
            if ca != cb:
                ch.append((i, j, ca, cb))
    changed[s] = ch
    print(f"{s}: raw cell diffs vs HEAD xlsx = {len(ch)}")

# --- PL sheet analysis
hdr = W["손익분해PL"][0]
print("PL hdr:", hdr)
assert hdr == H["손익분해PL"][0]
ch = changed["손익분해PL"]
keycols = hdr.index("원보험사코드"), hdr.index("항목번호"), hdr.index("공시분기")
iv, iq = hdr.index("값"), hdr.index("값_당분기")
print("changed columns:", Counter(hdr[c[1]] for c in ch), "(sheet col idx)")
print("changed rows:", len({c[0] for c in ch}))
print("changed companies:", Counter(W["손익분해PL"][c[0]][0] for c in ch))
print("changed quarters:", Counter(W["손익분해PL"][c[0]][hdr.index("공시분기")] for c in ch))
print("changed items:", sorted(Counter(W["손익분해PL"][c[0]][hdr.index("항목번호")] for c in ch).items(), key=lambda t: str(t[0])))
print("types of NEW 값_당분기 in changed cells:", Counter(type(c[3]).__name__ for c in ch if c[1] == iq), "| OLD:", Counter(type(c[2]).__name__ for c in ch if c[1] == iq))
print("types of NEW 값 in changed cells:", Counter(type(c[3]).__name__ for c in ch if c[1] == iv), "| OLD:", Counter(type(c[2]).__name__ for c in ch if c[1] == iv))

mn = gj("HEAD", "PL_breakdown.json")
mo = gj("HEAD~1", "PL_breakdown.json")


def key(r):
    return (r["원보험사코드"], str(r["항목번호"]), r["공시분기"])


# master-level change set (HEAD~1 -> HEAD)
dn = {key(r): r for r in mn}
do = {key(r): r for r in mo}
print("master rows old/new:", len(mo), len(mn), "keys equal:", set(dn) == set(do), "dup keys new:", len(mn) - len(dn))
mchg = set()
for k in dn:
    for col in ("값", "값_당분기"):
        if n(dn[k].get(col)) != n(do.get(k, {}).get(col)):
            mchg.add((k, col))
print("master cell changes HEAD~1->HEAD:", len(mchg), Counter(c for _, c in mchg), Counter(k[0] for k, _ in mchg))

# xlsx changed cell set
xchg = set()
bad_val = []
rows = W["손익분해PL"]
for c in ch:
    r = rows[c[0]]
    k = (r[0], str(r[hdr.index("항목번호")]), r[hdr.index("공시분기")])
    col = hdr[c[1]]
    xchg.add((k, col))
    # new value must equal HEAD master; old value must equal HEAD~1 master
    mv_new = n(dn[k][col]) if k in dn else "NOKEY"
    mv_old = n(do[k][col]) if k in do else "NOKEY"
    if n(c[3]) != mv_new:
        bad_val.append(("new!=headmaster", k, col, c[3], mv_new))
    if n(c[2]) != mv_old:
        bad_val.append(("old!=prevmaster", k, col, c[2], mv_old))
print("xlsx changed cells:", len(xchg), "| equal to master change set:", xchg == mchg, "| only_x", len(xchg - mchg), "only_m", len(mchg - xchg))
print("value mismatches in changed cells:", len(bad_val))
for b in bad_val[:10]:
    print("  ", b)

# whole-sheet multiset equality of working xlsx vs HEAD master ignoring 항목번호 column (koreanre pre-existing)
ign = hdr.index("항목번호")
xs = Counter(tuple(n(c) for k, c in enumerate(r) if k != ign) for r in rows[1:] if any(c is not None for c in r))
ms = Counter(tuple(n(r.get(h)) for k, h in enumerate(hdr) if k != ign) for r in mn)
print("PL sheet vs HEAD master (ignoring 항목번호): only_in_xlsx", sum((xs - ms).values()), "only_in_master", sum((ms - xs).values()))
xs2 = Counter(tuple(n(c) for c in r) for r in rows[1:] if any(c is not None for c in r))
ms2 = Counter(tuple(n(r.get(h)) for h in hdr) for r in mn)
print("PL sheet vs HEAD master (all cols): only_in_xlsx", sum((xs2 - ms2).values()), "only_in_master", sum((ms2 - xs2).values()))
blank = Counter(r[0] for r in rows[1:] if r[ign] is None)
print("blank 항목번호 rows by company (pre-existing?):", blank)
hb = Counter(r[0] for r in H["손익분해PL"][1:] if r[ign] is None)
print("same in HEAD xlsx:", hb == blank)
# row order/keys unchanged
print("row key order unchanged vs HEAD xlsx:", [tuple(r[k] for k in keycols) for r in rows] == [tuple(r[k] for k in keycols) for r in H["손익분해PL"]])
