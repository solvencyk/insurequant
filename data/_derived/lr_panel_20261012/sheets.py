from openpyxl import load_workbook
wb = load_workbook("C:/Users/sangwook.cho/Desktop/insurequant/insurequant_master_tables.xlsx", read_only=True, data_only=True)
for ws in wb.worksheets: print(ws.title, ws.max_row)
